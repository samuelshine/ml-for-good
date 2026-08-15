"""Build the CIA 3 presentation as both .pptx and a self-contained .html deck.

    .venv/bin/python presentation/build_deck.py
"""

from __future__ import annotations

import base64
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

from src import config as C

OUT_PPTX = ROOT / "presentation" / "CIA3_deck.pptx"
OUT_HTML = ROOT / "presentation" / "deck.html"

NAVY = RGBColor(0x1B, 0x2A, 0x41)
BLUE = RGBColor(0x2E, 0x5E, 0xAA)
RED = RGBColor(0xD1, 0x49, 0x5B)
GREEN = RGBColor(0x00, 0x91, 0x6E)
GREY = RGBColor(0x5A, 0x64, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Inches(13.333), Inches(7.5)


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------
def load():
    meta = json.loads((C.METRICS_DIR / "run_metadata.json").read_text())
    comp = pd.read_csv(C.METRICS_DIR / "model_comparison.csv")
    sig = json.loads((C.METRICS_DIR / "significance.json").read_text())
    fair = pd.read_csv(C.METRICS_DIR / "fairness_summary.csv")
    leak = pd.read_csv(C.METRICS_DIR / "leakage_sensitivity.csv")
    audit = pd.read_csv(C.METRICS_DIR / "cleaning_audit.csv")
    return meta, comp, sig, fair, leak, audit


META, COMP, SIG, FAIR, LEAK, AUDIT = load()
BEST = META["best_model"]
BASE = "Logistic Regression (baseline)"
best_row = COMP[COMP["model"] == BEST].iloc[0]
base_row = COMP[COMP["model"] == BASE].iloc[0]
BOOT = SIG["bootstrap_auc"]


# ---------------------------------------------------------------------------
# PPTX helpers
# ---------------------------------------------------------------------------
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, colour):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = colour
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def text(slide, x, y, w, h, content, size=18, bold=False, colour=NAVY,
         align="left", spacing=1.15):
    from pptx.enum.text import PP_ALIGN

    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0

    lines = content.split("\n") if isinstance(content, str) else content
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.line_spacing = spacing
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                       "right": PP_ALIGN.RIGHT}[align]
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = colour
            run.font.name = "Calibri"
    return box


def header(slide, title, subtitle=""):
    rect(slide, 0, 0, W, Inches(1.05), NAVY)
    text(slide, Inches(0.55), Inches(0.16), Inches(12), Inches(0.5),
         title, size=27, bold=True, colour=WHITE)
    if subtitle:
        text(slide, Inches(0.55), Inches(0.63), Inches(12), Inches(0.35),
             subtitle, size=13, colour=RGBColor(0xB8, 0xC4, 0xD4))


def picture(slide, fname, x, y, w=None, h=None):
    path = C.FIGURES_DIR / fname
    if not path.exists():
        return None
    kw = {}
    if w:
        kw["width"] = w
    if h:
        kw["height"] = h
    return slide.shapes.add_picture(str(path), x, y, **kw)


def bullets(slide, x, y, w, items, size=16, gap=0.46):
    for i, item in enumerate(items):
        marker = "▪"
        colour = NAVY
        if item.startswith("!"):
            marker, colour, item = "▲", RED, item[1:]
        elif item.startswith("+"):
            marker, colour, item = "●", GREEN, item[1:]
        text(slide, x, y + Inches(gap * i), Inches(0.3), Inches(0.4),
             marker, size=size, colour=colour, bold=True)
        text(slide, x + Inches(0.32), y + Inches(gap * i), w, Inches(0.4),
             item, size=size, colour=NAVY)


def stat(slide, x, y, w, value, label, colour=BLUE):
    rect(slide, x, y, w, Inches(1.5), RGBColor(0xF2, 0xF5, 0xF9))
    rect(slide, x, y, Inches(0.07), Inches(1.5), colour)
    text(slide, x + Inches(0.25), y + Inches(0.16), w - Inches(0.4), Inches(0.6),
         value, size=30, bold=True, colour=colour)
    text(slide, x + Inches(0.25), y + Inches(0.84), w - Inches(0.4), Inches(0.5),
         label, size=11, colour=GREY)


def table(slide, df, x, y, w, h, size=11, highlight_row=None):
    rows, cols = df.shape[0] + 1, df.shape[1]
    shape = slide.shapes.add_table(rows, cols, x, y, w, h).table

    for j, col in enumerate(df.columns):
        cell = shape.cell(0, j)
        cell.text = str(col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)
                r.font.bold = True
                r.font.color.rgb = WHITE

    for i in range(df.shape[0]):
        for j in range(cols):
            cell = shape.cell(i + 1, j)
            cell.text = str(df.iloc[i, j])
            cell.fill.solid()
            if highlight_row is not None and i == highlight_row:
                cell.fill.fore_color.rgb = RGBColor(0xDF, 0xF0, 0xE6)
            else:
                cell.fill.fore_color.rgb = (
                    WHITE if i % 2 == 0 else RGBColor(0xF4, 0xF6, 0xF9)
                )
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
                    r.font.color.rgb = NAVY
                    r.font.bold = highlight_row is not None and i == highlight_row
    return shape


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------
def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1 — Title
    s = blank(prs)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, Inches(3.03), W, Inches(0.05), RED)
    text(s, Inches(1), Inches(1.65), Inches(11.3), Inches(1.2),
         "Who Gets the Call?", size=50, bold=True, colour=WHITE)
    text(s, Inches(1), Inches(2.42), Inches(11.3), Inches(0.6),
         "Predicting 30-day hospital readmission for diabetic patients",
         size=21, colour=RGBColor(0xB8, 0xC4, 0xD4))
    text(s, Inches(1), Inches(3.35), Inches(11.3), Inches(1.6), [
        "ML for Social Good Ensemble Challenge  ·  CIA 3  ·  Mission Health",
        "MCA 521-4 Machine Learning  ·  CHRIST (Deemed to be University)",
    ], size=15, colour=WHITE, spacing=1.5)
    text(s, Inches(1), Inches(5.3), Inches(11.3), Inches(1.4), [
        f"Best model: {BEST.split(' (')[0]}   ·   Test ROC-AUC {best_row['roc_auc']:.4f}"
        f"   ·   {best_row['lift_at_10pct']:.2f}× lift at 10% capacity",
        "Dataset: UCI Diabetes 130-US Hospitals (Strack et al., 2014)  ·  CC BY 4.0",
    ], size=13, colour=RGBColor(0x8E, 0x9E, 0xB5), spacing=1.6)

    # 2 — Problem
    s = blank(prs)
    header(s, "The problem", "Mission Health · under-resourced clinics, fixed follow-up capacity")
    stat(s, Inches(0.55), Inches(1.4), Inches(3.9), "1 in 11",
         "diabetic inpatients readmitted within 30 days", RED)
    stat(s, Inches(4.72), Inches(1.4), Inches(3.9), "$15,000",
         "cost of one avoidable readmission", RED)
    stat(s, Inches(8.89), Inches(1.4), Inches(3.9), "$75",
         "cost of one nurse follow-up call", GREEN)
    text(s, Inches(0.55), Inches(3.3), Inches(12.2), Inches(0.5),
         "The gap is not knowing that calls work. It is knowing who to call.",
         size=21, bold=True, colour=NAVY)
    bullets(s, Inches(0.55), Inches(4.05), Inches(12), [
        "Beneficiaries: diabetic inpatients (especially uninsured and high-utilising)",
        "Discharge nurses, who get a ranked worklist instead of an undifferentiated ward list",
        "Safety-net hospitals carrying fixed capacity and rising CMS penalties",
        "Prediction target: P(readmission < 30 days), evaluated at the moment of discharge",
        "+Impact metric: precision@10% and lift@10% — not accuracy, which is meaningless at a 9% base rate",
    ], size=15)

    # 3 — Data
    s = blank(prs)
    header(s, "The data", "101,766 encounters · 130 US hospitals · peer-reviewed at source")
    bullets(s, Inches(0.55), Inches(1.35), Inches(6), [
        "UCI Diabetes 130-US Hospitals, 1999–2008",
        "Published with Strack et al. (2014), BioMed Research International",
        "CC BY 4.0 · de-identified under HIPAA Safe Harbor",
        "50 columns: demographics, utilisation, 23 drugs, ICD-9 diagnoses",
        "Chosen because it genuinely contains every wrangling problem the",
        "   brief asks us to handle — not a pre-cleaned benchmark",
    ], size=14, gap=0.42)
    picture(s, "01_missingness.png", Inches(6.9), Inches(1.4), w=Inches(6))
    text(s, Inches(0.55), Inches(4.5), Inches(6), Inches(2), [
        "Four columns, four different policies:",
        "  weight 96.9% → drop the value, keep 'was weighed' as a feature",
        "  medical_specialty 49.1% → explicit Unknown level",
        "  payer_code 39.6% → Unknown level; proxies for uninsured",
        "  race 2.2% → Unknown; never imputed, it is a protected attribute",
    ], size=13, colour=GREY, spacing=1.35)

    # 4 — Cleaning / leakage
    s = blank(prs)
    header(s, "Cleaning was the hard part", "Every dropped row is logged with a reason")
    picture(s, "03_encounters_per_patient.png", Inches(0.5), Inches(1.35), w=Inches(6.2))
    bullets(s, Inches(7.0), Inches(1.5), Inches(5.9), [
        "!101,766 rows are only 71,518 patients — one appears 40 times",
        "Rows are not independent observations",
        "Kept each patient's first encounter (as Strack et al. did)",
        "Dropped 2,423 encounters ending in death or hospice —",
        "   a patient who died cannot be readmitted",
        "+Positive rate fell 11.16% → 8.98% after de-duplication",
    ], size=14, gap=0.44)
    rect(s, Inches(7.0), Inches(5.05), Inches(5.9), Inches(1.75),
         RGBColor(0xFF, 0xF3, 0xF4))
    text(s, Inches(7.25), Inches(5.2), Inches(5.5), Inches(1.5), [
        "That 2.2-point drop is the finding:",
        "repeat encounters skew positive, so encounter-level",
        "modelling quietly flatters itself.",
    ], size=13, colour=RED, spacing=1.35)

    # 5 — EDA
    s = blank(prs)
    header(s, "What the data says", "Every claim is a plot, not an assertion")
    picture(s, "06_prior_inpatient.png", Inches(0.5), Inches(1.3), w=Inches(6.1))
    picture(s, "07_a1c_replication.png", Inches(6.85), Inches(1.55), w=Inches(6.0))
    bullets(s, Inches(0.55), Inches(5.35), Inches(12.2), [
        "Readmission rises monotonically with prior admissions: 8.1% → 29.3%",
        "+We reproduce the source paper's HbA1c finding (8.40% tested vs 9.11% not) — a credibility check on our cleaning",
        "HbA1c testing is one of the few actionable levers; age and race are not",
    ], size=14, gap=0.44)

    # 6 — Pipeline
    s = blank(prs)
    header(s, "Leakage-safe by construction", "Asserted in code, not claimed in prose")
    bullets(s, Inches(0.55), Inches(1.4), Inches(6.1), [
        "Stratified 60/20/20 split, seed 42",
        "One row per patient ⇒ no patient can straddle a split",
        "Every transform lives inside an sklearn Pipeline,",
        "   so CV refits it on each fold",
        "Test set untouched until the final comparison",
        "Class weights, not SMOTE: interpolating between one-hot",
        "   medical categories invents patients who cannot exist",
    ], size=14, gap=0.44)
    df = pd.DataFrame({
        "Split": ["Train", "Validation", "Test"],
        "Rows": ["41,991", "13,998", "13,998"],
        "Positive rate": ["8.98%", "8.98%", "8.98%"],
    })
    table(s, df, Inches(7.0), Inches(1.5), Inches(5.8), Inches(1.5), size=13)
    rect(s, Inches(7.0), Inches(3.5), Inches(5.8), Inches(3.1),
         RGBColor(0xF2, 0xF5, 0xF9))
    text(s, Inches(7.25), Inches(3.68), Inches(5.4), Inches(0.4),
         "Imputer chosen by experiment", size=16, bold=True, colour=NAVY)
    text(s, Inches(7.25), Inches(4.15), Inches(5.4), Inches(2.3), [
        "No numeric gaps remain after our policy, so we injected",
        "MCAR missingness to answer a deployment question:",
        "",
        "At 30% missing —  KNN 0.6295",
        "                              median 0.6130",
        "                              MICE 0.5992",
        "",
        "KNN degrades most gracefully.",
    ], size=12, colour=GREY, spacing=1.3)

    # 7 — Models
    s = blank(prs)
    header(s, "Baseline vs bagging vs boosting vs stacking",
           "Eight models, one untouched test set")
    picture(s, "11_model_comparison.png", Inches(0.4), Inches(1.25), w=Inches(12.5))
    text(s, Inches(0.55), Inches(6.35), Inches(12.2), Inches(0.8), [
        "Stacking uses 5-fold internal CV so base learners only ever produce out-of-fold "
        "predictions — that is what stops the",
        "meta-learner from rewarding whichever base model memorised hardest.",
    ], size=13, colour=GREY, spacing=1.3)

    # 8 — Does it beat the baseline
    s = blank(prs)
    header(s, "Does the ensemble beat the baseline?",
           "Answered with a confidence interval, not a bigger number")
    stat(s, Inches(0.55), Inches(1.35), Inches(3.9),
         f"{best_row['roc_auc']:.4f}", f"{BEST.split(' (')[0]} — test ROC-AUC", GREEN)
    stat(s, Inches(4.72), Inches(1.35), Inches(3.9),
         f"{base_row['roc_auc']:.4f}", "Logistic baseline — test ROC-AUC", GREY)
    stat(s, Inches(8.89), Inches(1.35), Inches(3.9),
         f"{BOOT['mean_difference']:+.4f}", "difference · 95% CI excludes zero", BLUE)
    picture(s, "17_bootstrap.png", Inches(1.4), Inches(3.15), w=Inches(10.5))
    text(s, Inches(0.55), Inches(5.6), Inches(12.2), Inches(1.3), [
        f"Yes — significantly, but modestly. Bootstrap 95% CI "
        f"[{BOOT['ci_low']:+.4f}, {BOOT['ci_high']:+.4f}] over 2,000 resamples; McNemar agrees.",
        "An AUC near 0.66 matches the published ceiling on this dataset. The limit is the data: "
        "housing, food security and",
        "caregiver support drive readmission, and none of them are in a discharge record. "
        "Reporting 0.95 here would mean a leak.",
    ], size=13, colour=NAVY, spacing=1.35)

    # 9 — Threshold
    s = blank(prs)
    header(s, "The threshold is a rationing decision",
           "0.5 is a library default, not a clinical choice")
    picture(s, "13_cost_threshold.png", Inches(0.5), Inches(1.3), w=Inches(7.3))
    bullets(s, Inches(8.1), Inches(1.5), Inches(4.8), [
        "A missed readmission costs ~200× a phone call",
        "!So unconstrained cost optimisation says: call everyone",
        "Correct for a clinic with unlimited nurses.",
        "   Useless for a real one.",
        "+Capacity — not cost — is the binding constraint",
        f"+Deployed threshold {META['threshold']:.3f}, flagging ~10%",
    ], size=14, gap=0.5)
    rect(s, Inches(8.1), Inches(5.0), Inches(4.8), Inches(1.75),
         RGBColor(0xE8, 0xF3, 0xEE))
    text(s, Inches(8.35), Inches(5.18), Inches(4.4), Inches(1.5), [
        "At 10% capacity:",
        f"  precision {best_row['precision_at_10pct']:.1%}  vs  9.0% base rate",
        f"  → {best_row['lift_at_10pct']:.2f}× more at-risk patients reached",
        "  with the same nurse hours.",
    ], size=13, colour=GREEN, spacing=1.35)

    # 10 — SHAP global
    s = blank(prs)
    header(s, "Why the model says what it says", "SHAP, cross-checked against two other methods")
    picture(s, "18_shap_global.png", Inches(0.45), Inches(1.3), w=Inches(6.3))
    bullets(s, Inches(7.1), Inches(1.5), Inches(5.8), [
        "Prior inpatient visits dominate — matching the EDA",
        "   gradient and mutual information, independently",
        "Directions match what we predicted before looking",
        "Cross-checked with permutation importance and LIME",
        "",
        "+Actionable: HbA1c testing, medication reconciliation,",
        "+   discharge destination",
        "Descriptive only: age, race, prior visit counts",
    ], size=14, gap=0.44)
    text(s, Inches(7.1), Inches(5.75), Inches(5.8), Inches(1.1),
         "A model whose top features were all descriptive would rank well "
         "and change nothing. Several of these are modifiable.",
         size=13, colour=GREY, spacing=1.3)

    # 11 — Local explanation
    s = blank(prs)
    header(s, "One patient, one explanation", "The live demo, on a synthetic record")
    pic = picture(s, "21_demo_patient.png", Inches(0.5), Inches(1.35), w=Inches(7.2))
    if pic is None:
        picture(s, "19_shap_local_true_positive.png", Inches(0.5), Inches(1.35), w=Inches(7.2))
    bullets(s, Inches(8.0), Inches(1.6), Inches(4.9), [
        "78-year-old, three prior admissions",
        "HbA1c above 8, no payer code recorded",
        "Discharged to a skilled nursing facility",
        "",
        "!Flagged for a follow-up call",
        "Each contribution shown in log-odds —",
        "   auditable, not decorative",
    ], size=14, gap=0.46)
    rect(s, Inches(8.0), Inches(5.35), Inches(4.9), Inches(1.4),
         RGBColor(0xFF, 0xF8, 0xE8))
    text(s, Inches(8.25), Inches(5.55), Inches(4.5), Inches(1.1), [
        "Abstention band 0.40–0.60 routes",
        "borderline cases to human review",
        "instead of auto-deciding.",
    ], size=13, colour=RGBColor(0x8A, 0x6D, 0x1F), spacing=1.35)

    # 12 — Fairness (the blocking finding)
    s = blank(prs)
    header(s, "The fairness audit failed", "And that is the most important slide here")
    fair_df = FAIR[["attribute", "disparate_impact_ratio", "passes_four_fifths_rule", "auc_gap"]].copy()
    fair_df.columns = ["Attribute", "Disparate impact", "Passes 4/5ths", "AUC gap"]
    fair_df["Disparate impact"] = fair_df["Disparate impact"].round(3)
    fair_df["AUC gap"] = fair_df["AUC gap"].round(3)
    fair_df["Passes 4/5ths"] = fair_df["Passes 4/5ths"].map({True: "Yes", False: "NO"})
    table(s, fair_df, Inches(0.55), Inches(1.35), Inches(6.1), Inches(2.0), size=12)

    rect(s, Inches(7.0), Inches(1.35), Inches(5.9), Inches(2.6),
         RGBColor(0xFF, 0xF3, 0xF4))
    text(s, Inches(7.25), Inches(1.55), Inches(5.5), Inches(2.3), [
        "Equal need, unequal service",
        "",
        "African American patients:  9.23% readmission rate",
        "Caucasian patients:              9.07% readmission rate",
        "",
        "Flagged for follow-up:  8.0%  vs  10.2%",
        "Recall:                            19.9%  vs  25.0%",
    ], size=13, colour=RED, spacing=1.35)

    bullets(s, Inches(0.55), Inches(4.2), Inches(12.2), [
        "!Not explained by differing risk — the underlying risk is the same",
        "!Label bias runs deeper: only same-network readmissions are recorded, so mobile and "
        "uninsured patients are",
        "!   systematically labelled negative. No fairness metric above can see this — they all "
        "use the same biased labels",
        "Per-group thresholds would equalise recall, but conditioning a clinical cut-off on race "
        "is not a fix we would deploy",
        "+This result only appeared once the threshold was set honestly. At a threshold that flags "
        "everyone, everything passes.",
    ], size=13, gap=0.42)

    # 13 — Limits
    s = blank(prs)
    header(s, "What this model cannot do", "Stated plainly, because the alternative is unsafe")
    bullets(s, Inches(0.55), Inches(1.4), Inches(12.2), [
        f"!Near-random (AUC 0.566) on repeat high-utilisers vs 0.687 on first admissions — "
        f"weakest where need is greatest",
        "!Fails the four-fifths rule on race, payer and age → not deployment-ready",
        "!Labels miss out-of-network readmissions, biased against the most vulnerable",
        "1999–2008 US data; not transferable to Indian clinics without prospective revalidation",
        "Triage aid only — never care denial, insurance pricing, discharge decisions, or staff evaluation",
        "Clinician override mandatory, and never logged against the clinician",
        "+Privacy: de-identified at source, no PII added, all demo records synthetic",
    ], size=15, gap=0.55)
    rect(s, Inches(0.55), Inches(5.75), Inches(12.2), Inches(1.05), RGBColor(0xF2, 0xF5, 0xF9))
    text(s, Inches(0.85), Inches(5.95), Inches(11.6), Inches(0.8),
         "The fairness audit was not a box to tick. It produced a result that blocks release — "
         "which is what an audit is for.",
         size=16, bold=True, colour=NAVY)

    # 14 — Close
    s = blank(prs)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, Inches(2.5), W, Inches(0.05), GREEN)
    text(s, Inches(1), Inches(1.3), Inches(11.3), Inches(1.0),
         "Same nurse hours.", size=42, bold=True, colour=WHITE)
    text(s, Inches(1), Inches(2.85), Inches(11.3), Inches(1.0),
         f"{best_row['lift_at_10pct']:.1f}× more at-risk patients reached.",
         size=42, bold=True, colour=GREEN)
    text(s, Inches(1), Inches(4.2), Inches(11.3), Inches(2), [
        "With an explanation attached to every flag,",
        "and an explicit statement of who the model fails.",
    ], size=19, colour=RGBColor(0xB8, 0xC4, 0xD4), spacing=1.5)
    text(s, Inches(1), Inches(6.2), Inches(11.3), Inches(1), [
        "make setup && make data && make train && make app",
        "Strack et al. (2014) · UCI ML Repository #296 · CC BY 4.0",
    ], size=13, colour=RGBColor(0x8E, 0x9E, 0xB5), spacing=1.5)

    prs.save(OUT_PPTX)
    print(f"Wrote {OUT_PPTX.relative_to(ROOT)} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


# ---------------------------------------------------------------------------
# HTML deck
# ---------------------------------------------------------------------------
def b64(fname: str) -> str:
    path = C.FIGURES_DIR / fname
    if not path.exists():
        return ""
    return "data:image/png;base64," + base64.b64encode(path.read_bytes()).decode()


def build_html() -> None:
    fair_rows = "".join(
        f"<tr><td>{r['attribute']}</td><td>{r['disparate_impact_ratio']:.3f}</td>"
        f"<td class='{'bad' if not r['passes_four_fifths_rule'] else 'good'}'>"
        f"{'NO' if not r['passes_four_fifths_rule'] else 'Yes'}</td>"
        f"<td>{r['auc_gap']:.3f}</td></tr>"
        for _, r in FAIR.iterrows()
    )
    model_rows = "".join(
        f"<tr class='{'hl' if r['model'] == BEST else ''}'><td>{r['model']}</td>"
        f"<td>{r['roc_auc']:.4f}</td><td>{r['pr_auc']:.4f}</td><td>{r['f1']:.4f}</td>"
        f"<td>{r['recall']:.4f}</td><td>{r['lift_at_10pct']:.2f}×</td></tr>"
        for _, r in COMP.iterrows()
    )

    slides = [
        f"""<section class="title">
<h1>Who Gets the Call?</h1>
<p class="sub">Predicting 30-day hospital readmission for diabetic patients</p>
<p class="meta">ML for Social Good Ensemble Challenge · CIA 3 · Mission Health<br>
MCA 521-4 Machine Learning · CHRIST (Deemed to be University)</p>
<p class="meta small">Best model {BEST.split(' (')[0]} · Test ROC-AUC {best_row['roc_auc']:.4f}
· {best_row['lift_at_10pct']:.2f}× lift at 10% capacity</p></section>""",

        f"""<section><h2>The problem</h2>
<div class="stats">
<div class="stat red"><b>1 in 11</b><span>diabetic inpatients readmitted within 30 days</span></div>
<div class="stat red"><b>$15,000</b><span>cost of one avoidable readmission</span></div>
<div class="stat green"><b>$75</b><span>cost of one nurse follow-up call</span></div>
</div>
<p class="lead">The gap is not knowing that calls work. It is knowing who to call.</p>
<ul>
<li>Beneficiaries: diabetic inpatients, discharge nurses, safety-net hospitals</li>
<li>Target: P(readmission &lt; 30 days), evaluated at discharge</li>
<li>Impact measured as precision@10% and lift@10% — not accuracy, which is
meaningless at a 9% base rate</li>
</ul></section>""",

        f"""<section><h2>Cleaning was the hard part</h2>
<img src="{b64('03_encounters_per_patient.png')}" alt="Encounters per patient">
<ul>
<li class="warn">101,766 rows are only 71,518 patients — one appears 40 times</li>
<li>Kept each patient's first encounter, as Strack et al. did</li>
<li>Dropped 2,423 encounters ending in death or hospice — they cannot be readmitted</li>
<li class="ok">Positive rate fell 11.16% → 8.98%: repeat encounters skew positive,
so encounter-level modelling flatters itself</li>
</ul></section>""",

        f"""<section><h2>What the data says</h2>
<div class="two"><img src="{b64('06_prior_inpatient.png')}" alt="Prior inpatient gradient">
<img src="{b64('07_a1c_replication.png')}" alt="HbA1c replication"></div>
<ul><li>Readmission rises monotonically with prior admissions: 8.1% → 29.3%</li>
<li class="ok">We reproduce the source paper's HbA1c finding — a credibility check on our cleaning</li></ul>
</section>""",

        f"""<section><h2>Baseline vs bagging vs boosting vs stacking</h2>
<table><thead><tr><th>Model</th><th>ROC-AUC</th><th>PR-AUC</th><th>F1</th>
<th>Recall</th><th>Lift@10%</th></tr></thead><tbody>{model_rows}</tbody></table>
<p class="note">Stacking uses 5-fold internal CV, so base learners only ever produce
out-of-fold predictions — that is what prevents meta-leakage.</p></section>""",

        f"""<section><h2>Does the ensemble beat the baseline?</h2>
<img src="{b64('17_bootstrap.png')}" alt="Bootstrap CI">
<p class="lead">Yes — significantly, but modestly.</p>
<p>{BOOT['mean_difference']:+.4f} ROC-AUC, bootstrap 95% CI
[{BOOT['ci_low']:+.4f}, {BOOT['ci_high']:+.4f}] over 2,000 resamples. The interval excludes zero.</p>
<p class="note">An AUC near 0.66 matches the published ceiling on this dataset. The limit is the
data: housing, food security and caregiver support drive readmission and none are in a discharge
record. Reporting 0.95 here would mean a leak, not a breakthrough.</p></section>""",

        f"""<section><h2>The threshold is a rationing decision</h2>
<img src="{b64('13_cost_threshold.png')}" alt="Cost threshold curve">
<ul><li class="warn">A missed readmission costs ~200× a call, so unconstrained
optimisation says "call everyone" — useless for a real clinic</li>
<li class="ok">Capacity, not cost, is the binding constraint</li>
<li class="ok">At 10% capacity: {best_row['precision_at_10pct']:.1%} precision vs 9.0% base rate
= {best_row['lift_at_10pct']:.2f}× more at-risk patients reached</li></ul></section>""",

        f"""<section><h2>Why the model says what it says</h2>
<img src="{b64('18_shap_global.png')}" alt="Global SHAP importance">
<ul><li>Prior inpatient visits dominate, matching the EDA and mutual information independently</li>
<li class="ok">Actionable levers: HbA1c testing, medication reconciliation, discharge destination</li>
<li>Descriptive only: age, race, prior visit counts</li></ul></section>""",

        f"""<section class="alert"><h2>The fairness audit failed</h2>
<table><thead><tr><th>Attribute</th><th>Disparate impact</th><th>Passes 4/5ths</th>
<th>AUC gap</th></tr></thead><tbody>{fair_rows}</tbody></table>
<p class="lead red">Equal need, unequal service.</p>
<p>African American patients have a readmission rate of <b>9.23%</b> versus <b>9.07%</b> for
Caucasian patients — yet they are flagged for follow-up <b>8.0%</b> of the time versus
<b>10.2%</b>, with recall <b>19.9%</b> versus <b>25.0%</b>.</p>
<p class="note">Deeper still: only same-network readmissions are recorded, so mobile and uninsured
patients are systematically labelled negative. No metric above can detect this — they are all
computed against the same biased labels.</p></section>""",

        f"""<section><h2>What this model cannot do</h2>
<ul>
<li class="warn">Near-random (AUC 0.566) on repeat high-utilisers vs 0.687 on first
admissions — weakest where need is greatest</li>
<li class="warn">Fails the four-fifths rule on race, payer and age → not deployment-ready</li>
<li>1999–2008 US data; not transferable to Indian clinics without prospective revalidation</li>
<li>Triage aid only — never care denial, insurance pricing, or discharge decisions</li>
<li class="ok">De-identified at source, no PII added, all demo records synthetic</li>
</ul>
<p class="lead">The fairness audit was not a box to tick. It produced a result that blocks
release — which is what an audit is for.</p></section>""",

        f"""<section class="title green-bg">
<h1>Same nurse hours.<br><span class="hi">{best_row['lift_at_10pct']:.1f}× more at-risk patients reached.</span></h1>
<p class="sub">With an explanation attached to every flag, and an explicit statement of who the
model fails.</p>
<p class="meta small"><code>make setup &amp;&amp; make data &amp;&amp; make train &amp;&amp; make app</code><br>
Strack et al. (2014) · UCI ML Repository #296 · CC BY 4.0</p></section>""",
    ]

    html = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Who Gets the Call? — CIA 3 Deck</title>
<style>
*{{box-sizing:border-box;margin:0;padding:0}}
body{{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif;
background:#0f1620;color:#1b2a41;line-height:1.55}}
section{{background:#fff;max-width:1100px;margin:26px auto;padding:44px 52px;
border-radius:14px;min-height:78vh;display:flex;flex-direction:column;justify-content:center;
box-shadow:0 8px 34px rgba(0,0,0,.34)}}
h1{{font-size:2.9rem;line-height:1.15;margin-bottom:18px}}
h2{{font-size:2rem;color:#1b2a41;margin-bottom:22px;padding-bottom:12px;
border-bottom:3px solid #2e5eaa}}
p{{margin-bottom:12px}}
.lead{{font-size:1.32rem;font-weight:600;margin:20px 0}}
.lead.red{{color:#d1495b}}
.note{{color:#5a6470;font-size:.95rem;margin-top:14px}}
.small{{font-size:.9rem}}
ul{{margin:14px 0 0 22px}}
li{{margin-bottom:9px}}
li.warn{{color:#d1495b;font-weight:600;list-style:none;position:relative}}
li.warn::before{{content:"▲";position:absolute;left:-22px}}
li.ok{{color:#00916e;font-weight:600;list-style:none;position:relative}}
li.ok::before{{content:"●";position:absolute;left:-22px}}
img{{max-width:100%;height:auto;display:block;margin:16px auto;border-radius:8px}}
.two{{display:grid;grid-template-columns:1fr 1fr;gap:14px;align-items:center}}
.stats{{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:18px;margin:22px 0}}
.stat{{background:#f2f5f9;border-left:6px solid #2e5eaa;padding:18px 20px;border-radius:8px}}
.stat.red{{border-color:#d1495b}} .stat.green{{border-color:#00916e}}
.stat b{{display:block;font-size:2.1rem;color:#2e5eaa;margin-bottom:6px}}
.stat.red b{{color:#d1495b}} .stat.green b{{color:#00916e}}
.stat span{{font-size:.9rem;color:#5a6470}}
table{{width:100%;border-collapse:collapse;margin:16px 0;font-size:.95rem;
display:block;overflow-x:auto;white-space:nowrap}}
th{{background:#1b2a41;color:#fff;padding:11px 13px;text-align:left}}
td{{padding:10px 13px;border-bottom:1px solid #e6eaf0}}
tr.hl td{{background:#dff0e6;font-weight:700}}
td.bad{{color:#d1495b;font-weight:700}} td.good{{color:#00916e;font-weight:700}}
section.title{{background:#1b2a41;color:#fff;text-align:left}}
section.title h2,section.title h1{{color:#fff;border:none}}
section.title .sub{{font-size:1.35rem;color:#b8c4d4;margin-bottom:26px}}
section.title .meta{{color:#8e9eb5;font-size:1rem}}
section.title code{{background:rgba(255,255,255,.12);padding:4px 9px;border-radius:5px}}
section.green-bg .hi{{color:#4ade80}}
section.alert{{border-top:7px solid #d1495b}}
@media (prefers-color-scheme:dark){{
 body{{background:#070b11}}
 section{{background:#141c26;color:#e6eaf0}}
 h2{{color:#e6eaf0}}
 .stat{{background:#1c2632}} .stat span{{color:#9aa8b8}}
 .note{{color:#9aa8b8}}
 td{{border-color:#25303e}}
 tr.hl td{{background:#16352a}}
}}
@media print{{section{{page-break-after:always;box-shadow:none;min-height:auto;margin:0}}}}
@media (max-width:720px){{section{{padding:26px 20px;margin:12px}}.two{{grid-template-columns:1fr}}
h1{{font-size:2rem}}h2{{font-size:1.5rem}}}}
</style></head><body>
{"".join(slides)}
</body></html>"""

    OUT_HTML.write_text(html, encoding="utf-8")
    size_mb = OUT_HTML.stat().st_size / 1e6
    print(f"Wrote {OUT_HTML.relative_to(ROOT)} ({len(slides)} slides, {size_mb:.1f} MB self-contained)")


if __name__ == "__main__":
    build_pptx()
    build_html()
